"""Ingestion pipeline for ERP and bank files.

The pipeline keeps the original payload in staging, normalizes rows into the
financial model, and materializes dashboard KPIs from the normalized data.
"""

import csv
import io
import json
import re
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from ..infrastructure.database import get_db_connection
from .forecast_service import calculate_daily_forecast


FIELD_ALIASES = {
    "transaction_date": ("transaction_date", "date", "fecha", "fecha_transaccion", "fecha movimiento"),
    "amount": ("amount", "monto", "importe", "valor", "amount_mxn"),
    "account_code": ("account_code", "cuenta", "codigo_cuenta", "account"),
    "account_name": ("account_name", "nombre_cuenta", "cuenta_nombre"),
    "category": ("category", "categoria", "tipo", "clasificacion"),
    "cost_center_code": ("cost_center_code", "centro_costo", "centro de costo", "cost_center"),
    "cost_center_name": ("cost_center_name", "nombre_centro", "centro_nombre"),
    "description": ("description", "descripcion", "concepto", "detalle"),
    "source_reference": ("source_reference", "referencia", "id_transaccion", "folio", "documento"),
    "currency": ("currency", "moneda", "divisa"),
}


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    text = str(value).strip().replace("$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as error:
        raise ValueError("Monto invalido: %s" % value) from error


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("Fecha invalida: %s" % value)


def _canonical_row(row: Dict[str, Any], position: int) -> Dict[str, Any]:
    normalized = {_key(key): value for key, value in row.items()}
    result: Dict[str, Any] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if _key(alias) in normalized and normalized[_key(alias)] not in (None, ""):
                result[field] = normalized[_key(alias)]
                break
    if "transaction_date" not in result or "amount" not in result:
        raise ValueError("La fila %s requiere fecha y monto" % position)
    result["transaction_date"] = _date(result["transaction_date"])
    result["amount"] = _decimal(result["amount"])
    result["account_code"] = str(result.get("account_code") or "GENERIC").strip()
    result["account_name"] = str(result.get("account_name") or result["account_code"]).strip()
    result["category"] = str(result.get("category") or "otros").strip().lower()
    result["cost_center_code"] = str(result.get("cost_center_code") or "GENERAL").strip()
    result["cost_center_name"] = str(result.get("cost_center_name") or result["cost_center_code"]).strip()
    result["description"] = str(result.get("description") or "Importado desde archivo").strip()
    result["source_reference"] = str(result.get("source_reference") or "fila-%s" % position).strip()
    result["currency"] = str(result.get("currency") or "MXN").strip().upper()[:3]
    return result


def extract_excel(content: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    return [dict(zip(headers, values)) for values in rows[1:] if any(value not in (None, "") for value in values)]


def extract_csv(content: bytes) -> List[Dict[str, Any]]:
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def extract_pdf(content: bytes) -> List[Dict[str, Any]]:
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    rows: List[Dict[str, Any]] = []
    pattern = re.compile(r"(?P<date>\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2}).*?(?P<amount>-?\$?\d[\d.,]*)")
    for line in text.splitlines():
        match = pattern.search(line)
        if match:
            rows.append({
                "fecha": match.group("date"),
                "monto": match.group("amount"),
                "descripcion": line[: match.start("amount")].strip(),
                "referencia": "pdf-%s" % len(rows),
            })
    if not rows:
        raise ValueError("No se encontraron filas financieras legibles en el PDF")
    return rows


def extract_raw_rows(file_name: str, content: bytes) -> List[Dict[str, Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        raw_rows = extract_excel(content)
    elif suffix == ".csv":
        raw_rows = extract_csv(content)
    elif suffix == ".pdf":
        raw_rows = extract_pdf(content)
    else:
        raise ValueError("Formato no soportado: %s" % suffix)
    return raw_rows


def extract_document_values(file_name: str, content: bytes) -> List[Dict[str, Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        pages = [
            {"page_number": page_number, "text": page.extract_text() or ""}
            for page_number, page in enumerate(PdfReader(io.BytesIO(content)).pages, start=1)
        ]
        if not any(page["text"].strip() for page in pages):
            raise ValueError("El PDF no contiene texto extraíble; los PDF escaneados aún no se pueden leer")
        return pages
    if suffix not in (".xlsx", ".xlsm", ".csv"):
        raise ValueError("Formato no soportado. Usa PDF, CSV o Excel (.xlsx, .xlsm)")

    rows = extract_raw_rows(file_name, content)
    if not rows:
        raise ValueError("El archivo no contiene valores para guardar")
    return [{"row_number": row_number, "values": row} for row_number, row in enumerate(rows, start=2)]


def extract_rows(file_name: str, content: bytes) -> List[Dict[str, Any]]:
    return [_canonical_row(row, index) for index, row in enumerate(extract_raw_rows(file_name, content), start=2)]


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, timedelta):
        return str(value)
    return value


def _source_id(cursor: Any, source_type: str, file_name: str) -> int:
    cursor.execute(
        "SELECT id FROM fuentes_datos WHERE nombre = %s AND tipo = %s::fuente_tipo ORDER BY id LIMIT 1",
        (file_name, source_type),
    )
    existing = cursor.fetchone()
    if existing is not None:
        return existing[0]
    cursor.execute(
        """
        INSERT INTO fuentes_datos (nombre, tipo, configuracion_conexion, esta_activa)
        VALUES (%s, %s::fuente_tipo, %s::jsonb, TRUE)
        RETURNING id
        """,
        (file_name, source_type, json.dumps({"mode": "file", "file_name": file_name})),
    )
    return cursor.fetchone()[0]


def _upsert_transaction(cursor: Any, row: Dict[str, Any], source_type: str, data_source_id: int) -> None:
    cursor.execute(
        """
        INSERT INTO cuentas (codigo, nombre, categoria)
        VALUES (%s, %s, %s)
        ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre, categoria = EXCLUDED.categoria
        RETURNING id
        """,
        (row["account_code"], row["account_name"], row["category"]),
    )
    account_id = cursor.fetchone()[0]
    cursor.execute(
        """
        INSERT INTO centros_costo (codigo, nombre)
        VALUES (%s, %s)
        ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre
        RETURNING id
        """,
        (row["cost_center_code"], row["cost_center_name"]),
    )
    cost_center_id = cursor.fetchone()[0]
    cursor.execute(
        """
        INSERT INTO transacciones
            (cuenta_id, centro_costo_id, fuente_datos_id, tipo_fuente, referencia_fuente, monto, moneda, fecha_transaccion, descripcion)
        VALUES (%s, %s, %s, %s::fuente_tipo, %s, %s, %s, %s, %s)
        ON CONFLICT (fuente_datos_id, referencia_fuente) DO UPDATE SET
            cuenta_id = EXCLUDED.cuenta_id,
            centro_costo_id = EXCLUDED.centro_costo_id,
            monto = EXCLUDED.monto,
            moneda = EXCLUDED.moneda,
            fecha_transaccion = EXCLUDED.fecha_transaccion,
            descripcion = EXCLUDED.descripcion
        """,
        (account_id, cost_center_id, data_source_id, source_type, row["source_reference"], row["amount"], row["currency"], row["transaction_date"], row["description"]),
    )


def _refresh_kpis(cursor: Any) -> None:
    cursor.execute(
        "UPDATE definiciones_kpi SET codigo = 'flujo_neto_bancario', nombre = 'Flujo neto bancario', descripcion_formula = 'Suma mensual de movimientos bancarios; no representa saldo sin saldo inicial' WHERE codigo = 'liquidez_inmediata'"
    )
    cursor.execute(
        "SELECT id, codigo FROM definiciones_kpi WHERE codigo IN ('flujo_neto_bancario', 'margen_operativo', 'ingresos_acumulados', 'desviacion_presupuestal')"
    )
    definition_ids = {code: definition_id for definition_id, code in cursor.fetchall()}
    cursor.execute(
        "DELETE FROM valores_kpi WHERE definicion_kpi_id IN (SELECT id FROM definiciones_kpi WHERE codigo IN ('flujo_neto_bancario', 'margen_operativo', 'ingresos_acumulados', 'desviacion_presupuestal')) AND centro_costo_id IS NULL"
    )

    periods: Dict[date, Dict[str, Decimal]] = {}
    cursor.execute(
        """
        SELECT date_trunc('month', fecha_transaccion)::date,
               COALESCE(SUM(CASE WHEN monto > 0 THEN monto ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN monto < 0 THEN ABS(monto) ELSE 0 END), 0)
        FROM transacciones
        GROUP BY 1
        ORDER BY 1
        """
    )
    for period, income, expenses in cursor.fetchall():
        periods[period] = {
            "income": Decimal(str(income)),
            "expenses": Decimal(str(expenses)),
            "has_actuals": True,
        }

    cursor.execute(
        """
        SELECT date_trunc('month', fecha_transaccion)::date, COALESCE(SUM(monto), 0)
        FROM transacciones
        WHERE tipo_fuente = 'banco'::fuente_tipo
        GROUP BY 1
        """
    )
    bank_flows = {period: Decimal(str(value)) for period, value in cursor.fetchall()}

    cursor.execute(
        """
        WITH latest_budgets AS (
            SELECT DISTINCT ON (b.cuenta_id, b.centro_costo_id, b.periodo)
                   b.periodo, b.monto_presupuestado
            FROM presupuestos b
            ORDER BY b.cuenta_id, b.centro_costo_id, b.periodo, b.version DESC
        )
        SELECT periodo, COALESCE(SUM(monto_presupuestado), 0)
        FROM latest_budgets
        GROUP BY periodo
        """
    )
    for period, budget in cursor.fetchall():
        periods.setdefault(period, {"income": Decimal("0"), "expenses": Decimal("0"), "has_actuals": False})["budget"] = Decimal(str(budget))

    now = datetime.now(timezone.utc)
    for period, values in periods.items():
        metrics: Dict[str, Decimal] = {}
        if period in bank_flows:
            metrics["flujo_neto_bancario"] = bank_flows[period]
        income = values["income"]
        if values["has_actuals"] and income > 0:
            metrics["margen_operativo"] = (income - values["expenses"]) / income * Decimal("100")
        if values["has_actuals"]:
            metrics["ingresos_acumulados"] = income
        budget = values.get("budget", Decimal("0"))
        if values["has_actuals"] and budget > 0:
            metrics["desviacion_presupuestal"] = (values["expenses"] - budget) / budget * Decimal("100")
        for code, value in metrics.items():
            definition_id = definition_ids.get(code)
            if definition_id is None:
                continue
            cursor.execute(
                """
                INSERT INTO valores_kpi (definicion_kpi_id, centro_costo_id, periodo, valor, calculado_en)
                VALUES (%s, NULL, %s, %s, %s)
                ON CONFLICT (definicion_kpi_id, centro_costo_id, periodo)
                DO UPDATE SET valor = EXCLUDED.valor, calculado_en = EXCLUDED.calculado_en
                """,
                (definition_id, period, value, now),
            )


def _refresh_alerts(cursor: Any) -> None:
    cursor.execute(
        """
        WITH latest_budgets AS (
            SELECT DISTINCT ON (b.cuenta_id, b.centro_costo_id, b.periodo)
                   b.cuenta_id, b.centro_costo_id, b.periodo, b.monto_presupuestado
            FROM presupuestos b
            ORDER BY b.cuenta_id, b.centro_costo_id, b.periodo, b.version DESC
        ), actuals AS (
            SELECT cuenta_id, centro_costo_id, date_trunc('month', fecha_transaccion)::date AS periodo,
                   ABS(SUM(monto)) AS monto_real
            FROM transacciones
            WHERE monto < 0
            GROUP BY cuenta_id, centro_costo_id, date_trunc('month', fecha_transaccion)::date
        )
        SELECT b.cuenta_id, b.centro_costo_id, b.periodo, b.monto_presupuestado, a.monto_real
        FROM latest_budgets b
        JOIN cuentas c ON c.id = b.cuenta_id
        LEFT JOIN actuals a ON a.cuenta_id = b.cuenta_id
            AND a.centro_costo_id IS NOT DISTINCT FROM b.centro_costo_id
            AND a.periodo = b.periodo
        WHERE lower(c.categoria) IN ('gasto', 'gastos', 'egreso', 'egresos', 'expense')
          AND b.monto_presupuestado > 0
                    AND a.cuenta_id IS NOT NULL
        """
    )
    for account_id, cost_center_id, period, budget, actual in cursor.fetchall():
        budget_amount = Decimal(str(budget))
        actual_amount = Decimal(str(actual))
        deviation = (actual_amount - budget_amount) / budget_amount * Decimal("100")
        if abs(deviation) < Decimal("10"):
            continue
        severity = "alta" if abs(deviation) >= Decimal("20") else "media"
        cursor.execute(
            """
            INSERT INTO alertas
                (cuenta_id, centro_costo_id, periodo, monto_presupuestado, monto_real, desviacion_pct, severidad, estado, detectado_en)
            VALUES (%s, %s, %s, %s, %s, %s, %s::severidad_alerta, 'nueva', %s)
            ON CONFLICT (periodo, centro_costo_id, cuenta_id) DO UPDATE SET
                monto_presupuestado = EXCLUDED.monto_presupuestado,
                monto_real = EXCLUDED.monto_real,
                desviacion_pct = EXCLUDED.desviacion_pct,
                severidad = EXCLUDED.severidad,
                detectado_en = EXCLUDED.detectado_en
            """,
            (account_id, cost_center_id, period, budget_amount, actual_amount, deviation, severity, datetime.now(timezone.utc)),
        )


def _refresh_bank_forecast(cursor: Any) -> None:
    cursor.execute(
        "SELECT MIN(fecha_transaccion), MAX(fecha_transaccion) FROM transacciones WHERE tipo_fuente = 'banco'::fuente_tipo"
    )
    first_day, last_day = cursor.fetchone()
    if first_day is None or last_day is None:
        return
    if (last_day - first_day).days + 1 < 180 or last_day < date.today() - timedelta(days=1):
        return

    training_start = max(first_day, last_day - timedelta(days=179))
    cursor.execute(
        """
        SELECT calendar.day::date, COALESCE(SUM(t.monto), 0)
        FROM generate_series(%s::date, %s::date, INTERVAL '1 day') AS calendar(day)
        LEFT JOIN transacciones t
          ON t.fecha_transaccion = calendar.day::date
         AND t.tipo_fuente = 'banco'::fuente_tipo
        GROUP BY calendar.day
        ORDER BY calendar.day
        """,
        (training_start, last_day),
    )
    daily_values = [(day, Decimal(str(amount))) for day, amount in cursor.fetchall()]
    forecast = calculate_daily_forecast(daily_values)
    if forecast is None:
        return

    generated_at = datetime.now(timezone.utc)
    notes = "Regresión lineal diaria; backtest holdout 30 días; MAE=%s MXN; intervalo predictivo 80%%" % forecast.backtest_mae
    cursor.execute(
        """
        INSERT INTO ejecuciones_pronostico
            (nombre_modelo, ejecutado_en, datos_entrenamiento_desde, datos_entrenamiento_hasta, notas)
        VALUES ('regresion_lineal_diaria_v1', %s, %s, %s, %s)
        RETURNING id
        """,
        (generated_at, training_start, last_day, notes),
    )
    forecast_run_id = cursor.fetchone()[0]
    for point in forecast.values:
        cursor.execute(
            """
            INSERT INTO valores_pronostico
                (ejecucion_pronostico_id, fecha_objetivo, valor_predicho, limite_inferior, limite_superior)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (forecast_run_id, point.target_date, point.predicted_value, point.lower_bound, point.upper_bound),
        )


def refresh_kpis_from_transactions() -> bool:
    """Refresh monthly financial metrics when actuals or budgets exist."""
    with get_db_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("SELECT EXISTS (SELECT 1 FROM transacciones LIMIT 1) OR EXISTS (SELECT 1 FROM presupuestos LIMIT 1)")
            has_financial_data = bool(cursor.fetchone()[0])
            if has_financial_data:
                _refresh_kpis(cursor)
                _refresh_alerts(cursor)
                conn.commit()
            return has_financial_data


def ingest_file(file_name: str, content: bytes, source_type: str = "erp") -> Dict[str, Any]:
    if source_type not in ("erp", "banco"):
        raise ValueError("source_type debe ser erp o banco")
    raw_rows = extract_raw_rows(file_name, content)
    if not raw_rows:
        raise ValueError("El archivo no contiene filas financieras")
    rows: List[Tuple[int, Dict[str, Any], Dict[str, Any]]] = []
    rejected_rows: List[Tuple[Dict[str, Any], str]] = []
    for index, raw_row in enumerate(raw_rows, start=2):
        try:
            row = _canonical_row(raw_row, index)
        except ValueError as error:
            rejected_rows.append((raw_row, str(error)))
            continue
        row["source_reference"] = "%s:%s" % (Path(file_name).name, row["source_reference"])
        rows.append((index, raw_row, row))
    if not rows and rejected_rows:
        raise ValueError("Todas las filas fueron rechazadas durante la normalización")

    started = datetime.now(timezone.utc)
    with get_db_connection() as conn:
        with conn.cursor() as cursor:
            data_source_id = _source_id(cursor, source_type, file_name)
            cursor.execute(
                "INSERT INTO registros_sincronizacion (fuente_datos_id, iniciado_en, estado, filas_ingresadas) VALUES (%s, %s, 'en_proceso', 0) RETURNING id",
                (data_source_id, started),
            )
            sync_log_id = cursor.fetchone()[0]
            staging_table = "datos_temporales_banco" if source_type == "banco" else "datos_temporales_erp"
            for row_number, raw_row, row in rows:
                payload = {
                    "file_name": Path(file_name).name,
                    "row_number": row_number,
                    "values": {key: _json_value(value) for key, value in raw_row.items()},
                }
                cursor.execute(
                    "INSERT INTO %s (registro_sincronizacion_id, carga_cruda, ingresado_en) VALUES (%%s, %%s::jsonb, %%s)" % staging_table,
                    (sync_log_id, json.dumps(payload, default=_json_value), started),
                )
                _upsert_transaction(cursor, row, source_type, data_source_id)
            for raw_row, reason in rejected_rows:
                payload = {key: _json_value(value) for key, value in raw_row.items()}
                cursor.execute(
                    """
                    INSERT INTO registros_sincronizacion_rechazados
                        (registro_sincronizacion_id, carga_cruda, motivo_rechazo)
                    VALUES (%s, %s::jsonb, %s)
                    """,
                    (sync_log_id, json.dumps(payload), reason),
                )
            _refresh_kpis(cursor)
            _refresh_alerts(cursor)
            if source_type == "banco":
                _refresh_bank_forecast(cursor)
            cursor.execute(
                """
                UPDATE registros_sincronizacion
                SET finalizado_en = %s, estado = 'exitoso', filas_ingresadas = %s, filas_rechazadas = %s
                WHERE id = %s
                """,
                (datetime.now(timezone.utc), len(rows), len(rejected_rows), sync_log_id),
            )
            conn.commit()
    return {
        "sync_log_id": sync_log_id,
        "data_source_id": data_source_id,
        "rows_ingested": len(rows),
        "rows_rejected": len(rejected_rows),
        "status": "exitoso",
    }


def _canonical_budget_row(row: Dict[str, Any], position: int) -> Dict[str, Any]:
    normalized = {_key(key): value for key, value in row.items()}

    def value_for(*aliases: str) -> Any:
        for alias in aliases:
            value = normalized.get(_key(alias))
            if value not in (None, ""):
                return value
        raise ValueError("La fila %s requiere periodo, monto presupuestado y cuenta" % position)

    period = _date(value_for("periodo", "period", "fecha", "date")).replace(day=1)
    amount = _decimal(value_for("monto_presupuestado", "presupuesto", "budget", "amount", "monto"))
    account_code = str(value_for("account_code", "cuenta", "codigo_cuenta", "account")).strip()
    account_name = str(normalized.get(_key("nombre_cuenta")) or normalized.get(_key("account_name")) or account_code).strip()
    cost_center_code = str(normalized.get(_key("centro_costo")) or normalized.get(_key("cost_center_code")) or "GENERAL").strip()
    cost_center_name = str(normalized.get(_key("nombre_centro")) or normalized.get(_key("cost_center_name")) or cost_center_code).strip()
    return {
        "period": period,
        "amount": amount,
        "account_code": account_code,
        "account_name": account_name,
        "cost_center_code": cost_center_code,
        "cost_center_name": cost_center_name,
        "category": "egreso",
    }


def ingest_budget_file(file_name: str, content: bytes) -> Dict[str, Any]:
    file_name = Path(file_name).name
    raw_rows = extract_raw_rows(file_name, content)
    if not raw_rows:
        raise ValueError("El archivo no contiene filas presupuestales")
    rows: List[Tuple[int, Dict[str, Any], Dict[str, Any]]] = []
    rejected_rows: List[Tuple[Dict[str, Any], str]] = []
    for row_number, raw_row in enumerate(raw_rows, start=2):
        try:
            rows.append((row_number, raw_row, _canonical_budget_row(raw_row, row_number)))
        except ValueError as error:
            rejected_rows.append((raw_row, str(error)))
    if not rows:
        raise ValueError("Todas las filas presupuestales fueron rechazadas")

    started = datetime.now(timezone.utc)
    with get_db_connection() as conn:
        with conn.cursor() as cursor:
            data_source_id = _source_id(cursor, "erp", file_name)
            cursor.execute(
                "INSERT INTO registros_sincronizacion (fuente_datos_id, iniciado_en, estado, filas_ingresadas) VALUES (%s, %s, 'en_proceso', 0) RETURNING id",
                (data_source_id, started),
            )
            sync_log_id = cursor.fetchone()[0]
            for row_number, raw_row, row in rows:
                cursor.execute(
                    "INSERT INTO cuentas (codigo, nombre, categoria) VALUES (%s, %s, %s) ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre RETURNING id",
                    (row["account_code"], row["account_name"], row["category"]),
                )
                account_id = cursor.fetchone()[0]
                cursor.execute(
                    "INSERT INTO centros_costo (codigo, nombre) VALUES (%s, %s) ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre RETURNING id",
                    (row["cost_center_code"], row["cost_center_name"]),
                )
                cost_center_id = cursor.fetchone()[0]
                cursor.execute(
                    "INSERT INTO presupuestos (cuenta_id, centro_costo_id, periodo, monto_presupuestado, version) VALUES (%s, %s, %s, %s, 1) ON CONFLICT (cuenta_id, centro_costo_id, periodo, version) DO UPDATE SET monto_presupuestado = EXCLUDED.monto_presupuestado",
                    (account_id, cost_center_id, row["period"], row["amount"]),
                )
                payload = {"file_name": file_name, "row_number": row_number, "values": raw_row}
                cursor.execute(
                    "INSERT INTO datos_temporales_erp (registro_sincronizacion_id, carga_cruda, ingresado_en) VALUES (%s, %s::jsonb, %s)",
                    (sync_log_id, json.dumps(payload, default=_json_value), started),
                )
            for raw_row, reason in rejected_rows:
                cursor.execute(
                    "INSERT INTO registros_sincronizacion_rechazados (registro_sincronizacion_id, carga_cruda, motivo_rechazo) VALUES (%s, %s::jsonb, %s)",
                    (sync_log_id, json.dumps({"file_name": file_name, "values": raw_row}, default=_json_value), reason),
                )
            _refresh_kpis(cursor)
            _refresh_alerts(cursor)
            cursor.execute(
                "UPDATE registros_sincronizacion SET finalizado_en = %s, estado = 'exitoso', filas_ingresadas = %s, filas_rechazadas = %s WHERE id = %s",
                (datetime.now(timezone.utc), len(rows), len(rejected_rows), sync_log_id),
            )
            conn.commit()
    return {
        "sync_log_id": sync_log_id,
        "data_source_id": data_source_id,
        "rows_ingested": len(rows),
        "rows_rejected": len(rejected_rows),
        "status": "exitoso",
    }


def ingest_erp_document(file_name: str, content: bytes) -> Dict[str, Any]:
    document_values = extract_document_values(file_name, content)
    file_name = Path(file_name).name
    started = datetime.now(timezone.utc)
    with get_db_connection() as conn:
        with conn.cursor() as cursor:
            data_source_id = _source_id(cursor, "erp", file_name)
            cursor.execute(
                "INSERT INTO registros_sincronizacion (fuente_datos_id, iniciado_en, estado, filas_ingresadas) VALUES (%s, %s, 'en_proceso', 0) RETURNING id",
                (data_source_id, started),
            )
            sync_log_id = cursor.fetchone()[0]
            for value in document_values:
                payload = {"file_name": file_name, **value}
                cursor.execute(
                    "INSERT INTO datos_temporales_erp (registro_sincronizacion_id, carga_cruda, ingresado_en) VALUES (%s, %s::jsonb, %s)",
                    (sync_log_id, json.dumps(payload, default=_json_value), started),
                )
            cursor.execute(
                "UPDATE registros_sincronizacion SET finalizado_en = %s, estado = 'exitoso', filas_ingresadas = %s, filas_rechazadas = 0 WHERE id = %s",
                (datetime.now(timezone.utc), len(document_values), sync_log_id),
            )
            conn.commit()
    return {
        "sync_log_id": sync_log_id,
        "data_source_id": data_source_id,
        "rows_ingested": len(document_values),
        "rows_rejected": 0,
        "status": "exitoso",
    }
